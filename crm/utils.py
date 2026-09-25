import os
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from django.conf import settings
from openpyxl import Workbook, load_workbook
from django.contrib.auth.models import User
from .models import Lead


def export_leads_xlsx(leads, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    headers = ["ID", "Name", "Company", "Email", "Phone", "Status", "Stage",
               "Source", "Owner", "Value", "Created By", "Updated By",
               "Created At", "Updated At", "Notes"]
    ws.append(["ID", "Name", "Company", "Email", "Phone",
           "Country", "City",                              # NEW
           "Status", "Stage", "Source", "Owner", "Value",
           "Created By", "Updated By", "Created At", "Updated At", "Notes"])

    for l in leads:
        ws.append([
            l.id, l.name, l.company, l.email, l.phone,
            l.country, l.city,                                 # NEW
            l.status, l.stage, l.source,
            l.owner.username if l.owner else "",
            float(l.value or 0),
            l.created_by.username if l.created_by else "",
            l.updated_by.username if l.updated_by else "",
            l.created_at.strftime("%Y-%m-%d %H:%M") if l.created_at else "",
            l.updated_at.strftime("%Y-%m-%d %H:%M") if l.updated_at else "",
            l.notes,
        ])
    wb.save(path)


def import_leads_xlsx(file_obj, user):
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0, 0
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    created = 0
    skipped = 0
    for row in rows[1:]:
        if not row or not row[idx.get("name", 0)]:
            continue
        try:
            name = str(row[idx["name"]]).strip()
            if not name:
                skipped += 1
                continue
            owner_username = row[idx["owner_username"]] if "owner_username" in idx else None
            owner = None
            if owner_username:
                owner = User.objects.filter(username=str(owner_username).strip()).first()
            Lead.objects.create(
                name=name,
                company=str(row[idx["company"]] or "") if "company" in idx else "",
                email=str(row[idx["email"]] or "") if "email" in idx else "",
                phone=str(row[idx["phone"]] or "") if "phone" in idx else "",
                country=str(row[idx["country"]] or "") if "country" in idx else "",   # NEW
                city=str(row[idx["city"]] or "") if "city" in idx else "",           # NEW
                stage=str(row[idx["stage"]] or "") if "stage" in idx else "",
                source=str(row[idx["source"]] or "") if "source" in idx else "",
                value=float(row[idx["value"]] or 0)
                    if "value" in idx and row[idx["value"]] not in (None, "") else 0,
                owner=owner,
                created_by=user,
                updated_by=user,
            )
            created += 1
        except Exception:
            skipped += 1
    return created, skipped


def create_backup() -> Path:
    """Zip the sqlite db and uploads into /data/backups/."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = settings.DATA_DIR / "backups" / f"backup_{timestamp}.zip"
    db_path = settings.DATA_DIR / "database.sqlite3"
    uploads = settings.DATA_DIR / "uploads"

    with zipfile.ZipFile(backup_path, "w", zipfile.ZIP_DEFLATED) as zf:
        if db_path.exists():
            zf.write(db_path, arcname="database.sqlite3")
        if uploads.exists():
            for root, _, files in os.walk(uploads):
                for f in files:
                    full = Path(root) / f
                    rel = full.relative_to(uploads)
                    zf.write(full, arcname=f"uploads/{rel}")
    return backup_path